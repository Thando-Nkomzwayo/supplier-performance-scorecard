import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

# Create Workbook
wb = openpyxl.Workbook()

# Setup Sheets
ws_summary = wb.active
ws_summary.title = "Supplier Scorecard"
ws_data = wb.create_sheet(title="PO Data Entry")

# Enable Gridlines
ws_summary.views.sheetView[0].showGridLines = True
ws_data.views.sheetView[0].showGridLines = True

# --- Color Palette & Styling ---
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

TITLE_FONT = Font(name="Segoe UI", size=16, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="595959")
SECTION_FONT = Font(name="Segoe UI", size=12, bold=True, color="1F4E78")

BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
REGULAR_FONT = Font(name="Segoe UI", size=10)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

TOTAL_BORDER = Border(
    top=Side(style='thin', color='1F4E78'),
    bottom=Side(style='double', color='1F4E78')
)

# --- 1. Populate PO Data Entry Sheet ---
ws_data['A1'] = "PROCUREMENT PURCHASE ORDER TRACKER"
ws_data['A1'].font = TITLE_FONT
ws_data['A2'] = "Enter purchase order delivery performance details below. Excel formulas automatically compute variance and OTD status."
ws_data['A2'].font = SUBTITLE_FONT

po_headers = ["PO Number", "Supplier Name", "Order Date", "Promised Days", "Actual Days", "Lead Time Variance", "On-Time Status", "Order Value (ZAR)"]
for col_num, header in enumerate(po_headers, 1):
    cell = ws_data.cell(row=4, column=col_num, value=header)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

sample_po_data = [
    ["PO-1001", "Apex Logistics", "2026-10-01", 5, 5, "=E5-D5", '=IF(F5<=0, "On-Time", "Delayed")', 15000.00],
    ["PO-1002", "Apex Logistics", "2026-10-03", 5, 6, "=E6-D6", '=IF(F6<=0, "On-Time", "Delayed")', 22000.00],
    ["PO-1003", "Beta Industrial", "2026-10-05", 7, 7, "=E7-D7", '=IF(F7<=0, "On-Time", "Delayed")', 8500.00],
    ["PO-1004", "Beta Industrial", "2026-10-08", 7, 9, "=E8-D8", '=IF(F8<=0, "On-Time", "Delayed")', 12000.00],
    ["PO-1005", "Delta Steel", "2026-10-10", 4, 3, "=E9-D9", '=IF(F9<=0, "On-Time", "Delayed")', 31000.00],
    ["PO-1006", "Delta Steel", "2026-10-12", 4, 4, "=E10-D10", '=IF(F10<=0, "On-Time", "Delayed")', 18000.00],
    ["PO-1007", "Apex Logistics", "2026-10-14", 5, 4, "=E11-D11", '=IF(F11<=0, "On-Time", "Delayed")', 19500.00],
    ["PO-1008", "Beta Industrial", "2026-10-15", 7, 8, "=E12-D12", '=IF(F12<=0, "On-Time", "Delayed")', 14200.00],
]

for row_idx, row_data in enumerate(sample_po_data, 5):
    for col_idx, val in enumerate(row_data, 1):
        cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if col_idx in [4, 5, 6]:
            cell.alignment = Alignment(horizontal="center")
            cell.number_format = '#,##0'
        elif col_idx == 7:
            cell.alignment = Alignment(horizontal="center")
        elif col_idx == 8:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = 'R #,##0.00'
        elif col_idx in [1, 3]:
            cell.alignment = Alignment(horizontal="center")

# --- 2. Populate Supplier Scorecard Summary Sheet ---
ws_summary['A1'] = "AUTOMATED SUPPLIER PERFORMANCE SCORECARD"
ws_summary['A1'].font = TITLE_FONT
ws_summary['A2'] = "Automated summary aggregating vendor delivery performance, spend, and risk ratings."
ws_summary['A2'].font = SUBTITLE_FONT

scorecard_headers = ["Supplier Name", "Total Orders", "On-Time Orders", "On-Time Delivery (OTD %)", "Average Lead Time (Days)", "Total Spend (ZAR)", "Risk Rating"]
for col_num, header in enumerate(scorecard_headers, 1):
    cell = ws_summary.cell(row=4, column=col_num, value=header)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

suppliers = ["Apex Logistics", "Beta Industrial", "Delta Steel"]

for idx, supp in enumerate(suppliers, 5):
    ws_summary.cell(row=idx, column=1, value=supp).font = BOLD_FONT
    ws_summary.cell(row=idx, column=1).border = THIN_BORDER
    
    # Formulas linking to PO Data Entry sheet
    ws_summary.cell(row=idx, column=2, value=f'=COUNTIF(\'PO Data Entry\'!$B$5:$B$12, A{idx})').alignment = Alignment(horizontal="center")
    ws_summary.cell(row=idx, column=3, value=f'=COUNTIFS(\'PO Data Entry\'!$B$5:$B$12, A{idx}, \'PO Data Entry\'!$G$5:$G$12, "On-Time")').alignment = Alignment(horizontal="center")
    
    otd_cell = ws_summary.cell(row=idx, column=4, value=f'=C{idx}/B{idx}')
    otd_cell.alignment = Alignment(horizontal="center")
    otd_cell.number_format = '0.0%'
    
    avg_lt_cell = ws_summary.cell(row=idx, column=5, value=f'=AVERAGEIF(\'PO Data Entry\'!$B$5:$B$12, A{idx}, \'PO Data Entry\'!$E$5:$E$12)')
    avg_lt_cell.alignment = Alignment(horizontal="center")
    avg_lt_cell.number_format = '0.0'
    
    spend_cell = ws_summary.cell(row=idx, column=6, value=f'=SUMIF(\'PO Data Entry\'!$B$5:$B$12, A{idx}, \'PO Data Entry\'!$H$5:$H$12)')
    spend_cell.alignment = Alignment(horizontal="right")
    spend_cell.number_format = 'R #,##0.00'
    
    risk_cell = ws_summary.cell(row=idx, column=7, value=f'=IF(D{idx}>=0.85, "Low Risk", IF(D{idx}>=0.70, "Medium Risk", "High Risk"))')
    risk_cell.alignment = Alignment(horizontal="center")
    risk_cell.font = BOLD_FONT

    for col_c in range(1, 8):
        ws_summary.cell(row=idx, column=col_c).border = THIN_BORDER
        if col_c != 1 and col_c != 7:
            ws_summary.cell(row=idx, column=col_c).font = REGULAR_FONT

# Total Row on Summary Sheet
ws_summary.cell(row=8, column=1, value="Total / Portfolio Average").font = BOLD_FONT
ws_summary.cell(row=8, column=1).border = TOTAL_BORDER
ws_summary.cell(row=8, column=2, value='=SUM(B5:B7)').font = BOLD_FONT
ws_summary.cell(row=8, column=2).alignment = Alignment(horizontal="center")
ws_summary.cell(row=8, column=2).border = TOTAL_BORDER

ws_summary.cell(row=8, column=3, value='=SUM(C5:C7)').font = BOLD_FONT
ws_summary.cell(row=8, column=3).alignment = Alignment(horizontal="center")
ws_summary.cell(row=8, column=3).border = TOTAL_BORDER

total_otd = ws_summary.cell(row=8, column=4, value='=C8/B8')
total_otd.font = BOLD_FONT
total_otd.alignment = Alignment(horizontal="center")
total_otd.number_format = '0.0%'
total_otd.border = TOTAL_BORDER

total_avg_lt = ws_summary.cell(row=8, column=5, value='=AVERAGE(E5:E7)')
total_avg_lt.font = BOLD_FONT
total_avg_lt.alignment = Alignment(horizontal="center")
total_avg_lt.number_format = '0.0'
total_avg_lt.border = TOTAL_BORDER

total_spend = ws_summary.cell(row=8, column=6, value='=SUM(F5:F7)')
total_spend.font = BOLD_FONT
total_spend.alignment = Alignment(horizontal="right")
total_spend.number_format = 'R #,##0.00'
total_spend.border = TOTAL_BORDER

ws_summary.cell(row=8, column=7, value="—").alignment = Alignment(horizontal="center")
ws_summary.cell(row=8, column=7).font = BOLD_FONT
ws_summary.cell(row=8, column=7).border = TOTAL_BORDER

# --- Conditional Formatting for Risk Rating ---
green_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
green_font = Font(color="274E13", bold=True)
amber_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
amber_font = Font(color="7F6000", bold=True)
red_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
red_font = Font(color="783F04", bold=True)

ws_summary.conditional_formatting.add("G5:G7", CellIsRule(operator='equal', formula=['"Low Risk"'], fill=green_fill, font=green_font))
ws_summary.conditional_formatting.add("G5:G7", CellIsRule(operator='equal', formula=['"Medium Risk"'], fill=amber_fill, font=amber_font))
ws_summary.conditional_formatting.add("G5:G7", CellIsRule(operator='equal', formula=['"High Risk"'], fill=red_fill, font=red_font))

# --- Adjust Column Widths ---
for ws in [ws_summary, ws_data]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row in [1, 2]: # Skip titles for width calc
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

# Save Workbook
wb.save("Supplier_Performance_Scorecard.xlsx")